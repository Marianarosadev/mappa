from ._anvil_designer import tpvQueryTemplate
from anvil import *
import anvil.server
import anvil.tables as tables
import anvil.tables.query as q
from anvil.tables import app_tables
import anvil.users
import anvil.http
import openpyxl
from anvil import download

class tpvQuery(tpvQueryTemplate):
  def __init__(self, **properties):
    self.init_components(**properties)

  def button_1_click(self, **event_args):
    if self.date_picker_1.date is not None:
      response = anvil.server.call(
        'readTpv',
        self.date_picker_1.date,  
        self.date_picker_2.date, 
        self.text_box_1.text,
        self.text_box_2.text, 
      )
      print('response.data: ', response['data'])
      
      self.repeating_panel_1.items = response['data']
    else:
      self.label_5.text = 'Campo obrigatório'
      self.label_5.visible = True

  def button_2_click(self, **event_args):
    # Crie um novo arquivo XLSX
    workbook = openpyxl.Workbook()
    sheet = workbook.active

    data_grid = self.data_grid_1
    for row in data_grid.get_components():
      for col, cell in enumerate(row):
        sheet.cell(row=row.index + 1, column=col + 1, value=cell.text)

    # Salve o arquivo em um caminho temporário
    temp_file_path = "temp.xlsx"
    
    workbook.save(temp_file_path)
    
    # Inicie o download do arquivo
    download(temp_file_path)


